import csv
import sqlite3
from datetime import date, datetime
from pathlib import Path
import tkinter as tk
import tkinter.font as tkfont
from tkinter import filedialog, messagebox, ttk

APP_DIR = Path.home() / ".buh_calculator"
APP_DIR.mkdir(exist_ok=True)
DB_PATH = APP_DIR / "buh_calculator.db"

FACE = "#D4D0C8"
SHADOW = "#808080"
DARK = "#404040"
LIGHT = "#FFFFFF"
NAVY = "#0A246A"
MDI_BG = "#808080"
INFO = "#FFFFE1"
TITLE_FROM, TITLE_TO = "#F5C800", "#FFFFFF"
INACTIVE_FROM, INACTIVE_TO = "#808080", "#C0C0C0"
YELLOW_SEL = "#FFD84D"
YELLOW_SOFT = "#FFF3B0"
GREEN, RED, GRAY = "#006000", "#B00000", "#606060"
FONT_CANDIDATES = ["Tahoma", "Verdana", "MS Sans Serif", "Liberation Sans", "DejaVu Sans", "Noto Sans"]
UI_FAMILY = "TkDefaultFont"
LABEL_COL = 250

CATEGORIES = ["Продажи", "Услуги", "Зарплата", "Аренда", "Закупки", "Налоги", "Транспорт", "Прочее"]
VAT_RATES = ["0%", "5%", "10%", "20%"]

def money(value):
    return f"{value:,.2f}".replace(",", " ").replace(".", ",")

def blend(c1, c2, t):
    a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
    return "#" + "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(a, b))

class Database:
    def __init__(self):
        with sqlite3.connect(DB_PATH) as con:
            con.execute("""CREATE TABLE IF NOT EXISTS entries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                entry_date TEXT NOT NULL,
                kind TEXT NOT NULL,
                category TEXT NOT NULL,
                description TEXT,
                amount REAL NOT NULL,
                vat_rate REAL NOT NULL DEFAULT 0
            )""")

    def add(self, entry_date, kind, category, description, amount, vat_rate):
        with sqlite3.connect(DB_PATH) as con:
            con.execute("INSERT INTO entries(entry_date,kind,category,description,amount,vat_rate) VALUES(?,?,?,?,?,?)",
                        (entry_date, kind, category, description, amount, vat_rate))

    def all(self, start=None, end=None):
        sql = "SELECT id, entry_date, kind, category, description, amount, vat_rate FROM entries"
        args = []
        if start and end:
            sql += " WHERE entry_date BETWEEN ? AND ?"
            args = [start, end]
        sql += " ORDER BY entry_date DESC, id DESC"
        with sqlite3.connect(DB_PATH) as con:
            return con.execute(sql, args).fetchall()

    def delete(self, entry_id):
        with sqlite3.connect(DB_PATH) as con:
            con.execute("DELETE FROM entries WHERE id=?", (entry_id,))

def draw_icon(c, kind, x, y):
    def R(x1, y1, w, h, fill, outline=None):
        c.create_rectangle(x + x1, y + y1, x + x1 + w, y + y1 + h, fill=fill,
                           outline=outline or "", width=1 if outline else 0)

    def P(points, fill, outline="#000000"):
        c.create_polygon(*[v + (x if i % 2 == 0 else y) for i, v in enumerate(points)], fill=fill, outline=outline)

    def T(tx, ty, s, color):
        c.create_text(x + tx, y + ty, text=s, fill=color, font=(UI_FAMILY, 7, "bold"))

    def doc():
        P([3, 1, 10, 1, 13, 4, 13, 15, 3, 15], LIGHT)

    if kind == "app":
        R(1, 1, 14, 14, "#C00000"); T(8, 8, "Б", "#FFD800")
    elif kind == "overview":
        R(1, 14, 14, 1, "#000000"); R(2, 8, 3, 6, "#008000"); R(6, 3, 3, 11, NAVY); R(10, 6, 3, 8, "#C00000")
    elif kind == "journal":
        R(2, 1, 12, 14, "#000080"); R(3, 2, 10, 12, "#3A6EA5")
        for k in range(3):
            R(5, 4 + k * 3, 6, 1, LIGHT)
    elif kind == "calc":
        R(2, 1, 12, 14, DARK); R(3, 2, 10, 4, "#C8E8C8")
        for i in range(3):
            for j in range(3):
                R(3 + i * 4, 7 + j * 3, 3, 2, "#E0E0E0")
    elif kind == "usn":
        doc(); T(8, 9, "%", "#000080")
    elif kind == "sick":
        doc(); R(7, 5, 2, 8, "#E00000"); R(4, 8, 8, 2, "#E00000")
    elif kind == "maternity":
        doc(); T(8, 9, "♥", "#D0306A")
    elif kind == "vacation":
        for dx, dy in [(0, -7), (0, 6), (-7, 0), (6, 0), (-5, -5), (4, 4), (-5, 4), (4, -5)]:
            R(8 + dx, 8 + dy, 1, 1, "#E09000")
        c.create_oval(x + 4, y + 4, x + 12, y + 12, fill="#FFD800", outline="#B08000")
    elif kind == "new":
        doc(); R(9, 0, 2, 6, "#E0B000"); R(7, 2, 6, 2, "#E0B000")
    elif kind == "delete":
        c.create_line(x + 3, y + 3, x + 13, y + 13, fill="#C00000", width=2)
        c.create_line(x + 13, y + 3, x + 3, y + 13, fill="#C00000", width=2)
    elif kind == "refresh":
        c.create_arc(x + 3, y + 3, x + 13, y + 13, start=40, extent=280, style="arc", outline="#008000", width=2)
        P([9, 1, 14, 4, 10, 7], "#008000", "#008000")
    elif kind == "save":
        R(1, 1, 14, 14, "#000080"); R(4, 1, 8, 5, LIGHT); R(4, 9, 8, 6, "#C0C0C0")
    elif kind == "help":
        c.create_oval(x + 1, y + 1, x + 15, y + 15, fill=INFO, outline="#000080")
        T(8, 8, "?", "#000080")

class ChildWindow(tk.Frame):

    def __init__(self, app, key, title, icon, builder, resizable):
        super().__init__(app.mdi, bg=FACE, bd=2, relief="raised")
        self.app, self.key, self.title, self.icon = app, key, title, icon
        self.builder, self.resizable = builder, resizable
        self.active = False
        self.maximized = False
        self.pos = (0, 0)
        self.drag_offset = None
        self.calc = None
        self.bar_h = app.f_bold.metrics("linespace") + 6
        self.bar = tk.Canvas(self, height=self.bar_h, bg=TITLE_FROM, highlightthickness=0, bd=0)
        self.bar.pack(fill="x", padx=1, pady=1)
        self.body = tk.Frame(self, bg=FACE)
        self.body.pack(fill="both", expand=True, padx=10, pady=(8, 10))
        self.bar.bind("<Configure>", lambda e: self.draw_bar())
        self.bar.bind("<Button-1>", self.on_bar_click)
        self.bar.bind("<B1-Motion>", self.on_drag)
        self.bar.bind("<Double-Button-1>", lambda e: self.toggle_max())
        self.rebuild()

    def rebuild(self):
        for w in self.body.winfo_children():
            w.destroy()
        self.builder(self.body, self)

    def bar_buttons(self):
        w, bh = self.bar.winfo_width(), self.bar_h - 6
        bw = bh + 2
        buttons = [("close", w - bw - 3)]
        if self.resizable:
            buttons.append(("max", w - 2 * bw - 5))
        return [(kind, bx, 3, bw, bh) for kind, bx in buttons]

    def draw_bar(self):
        c = self.bar
        c.delete("all")
        w, h = c.winfo_width(), self.bar_h
        a, b = (TITLE_FROM, TITLE_TO) if self.active else (INACTIVE_FROM, INACTIVE_TO)
        for x in range(0, w, 2):
            c.create_rectangle(x, 0, x + 2, h, fill=blend(a, b, x / max(1, w)), width=0)
        draw_icon(c, self.icon, 3, (h - 16) // 2)
        c.create_text(24, h // 2, text=self.title, anchor="w", font=self.app.f_bold,
                      fill="#000000" if self.active else "#E8E8E8")
        for kind, x, y, bw, bh in self.bar_buttons():
            c.create_rectangle(x, y, x + bw, y + bh, fill=FACE, width=0)
            c.create_line(x, y + bh - 1, x, y, x + bw - 1, y, fill=LIGHT)
            c.create_line(x, y + bh - 1, x + bw - 1, y + bh - 1, x + bw - 1, y - 1, fill=DARK)
            c.create_line(x + 1, y + bh - 2, x + bw - 2, y + bh - 2, x + bw - 2, y, fill=SHADOW)
            if kind == "close":
                c.create_line(x + 4, y + 3, x + bw - 5, y + bh - 4, width=2)
                c.create_line(x + bw - 5, y + 3, x + 4, y + bh - 4, width=2)
            else:
                c.create_rectangle(x + 4, y + 3, x + bw - 5, y + bh - 4)
                c.create_line(x + 4, y + 4, x + bw - 5, y + 4, width=2)

    def on_bar_click(self, e):
        for kind, x, y, bw, bh in self.bar_buttons():
            if x <= e.x <= x + bw and y <= e.y <= y + bh:
                self.drag_offset = None
                self.close() if kind == "close" else self.toggle_max()
                return
        self.app.activate(self)
        self.drag_offset = (e.x_root - self.winfo_x(), e.y_root - self.winfo_y())

    def on_drag(self, e):
        if self.maximized or not self.drag_offset:
            return
        mdi = self.app.mdi
        x = min(max(e.x_root - self.drag_offset[0], 60 - self.winfo_width()), mdi.winfo_width() - 60)
        y = min(max(e.y_root - self.drag_offset[1], 0), mdi.winfo_height() - self.bar_h)
        self.move_to(x, y)

    def move_to(self, x, y):
        self.pos = (x, y)
        self.place(x=x, y=y)

    def toggle_max(self):
        if not self.resizable:
            return
        self.place_forget()
        if self.maximized:
            self.place(x=self.pos[0], y=self.pos[1])
        else:
            self.place(x=0, y=0, relwidth=1, relheight=1)
        self.maximized = not self.maximized

    def close(self):
        self.app.close_window(self)

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        global UI_FAMILY
        self.title("БухКальк — бухгалтерский калькулятор")
        self.geometry("1180x760")
        self.minsize(980, 650)
        self.configure(bg=FACE)
        families = set(tkfont.families(self))
        UI_FAMILY = next((f for f in FONT_CANDIDATES if f in families), "TkDefaultFont")
        self.f = tkfont.Font(self, family=UI_FAMILY, size=9)
        self.f_bold = tkfont.Font(self, family=UI_FAMILY, size=9, weight="bold")
        self.f_total = tkfont.Font(self, family=UI_FAMILY, size=11, weight="bold")
        self.db = Database()
        self.windows = {}
        self.active_win = None
        self.pages = {
            "overview": ("Обзор: сводка по всем операциям", "overview", self.page_overview, True),
            "entries": ("Журнал операций", "journal", self.page_entries, True),
            "new_entry": ("Операция: Новая", "new", self.page_new_entry, False),
            "calculator": ("Калькулятор НДС", "calc", self.page_calculator, False),
            "usn": ("Расчёт налога УСН", "usn", self.page_usn, False),
            "sick": ("Калькулятор больничного листа", "sick", self.page_sick, False),
            "maternity": ("Калькулятор декретных", "maternity", self.page_maternity, False),
            "vacation": ("Калькулятор отпускных", "vacation", self.page_vacation, False),
            "reports": ("Выгрузка отчёта", "save", self.page_reports, False),
        }
        self.setup_style()
        self.build_ui()
        self.bind_all("<Button-1>", self.on_global_click, add="+")
        self.bind("<Return>", lambda e: self.active_win and self.active_win.calc and self.active_win.calc())
        self.bind("<Escape>", lambda e: self.active_win and not self.active_win.resizable and self.active_win.close())
        for i, key in enumerate(["calculator", "usn", "sick", "maternity", "vacation"], start=1):
            self.bind(f"<Control-Key-{i}>", lambda e, k=key: self.show_page(k))
        self.bind("<Control-n>", lambda e: self.show_page("new_entry"))
        self.bind("<Control-j>", lambda e: self.show_page("entries"))
        self.bind("<F1>", lambda e: self.about())
        self.show_page("overview")

    def setup_style(self):
        self.option_add("*Font", self.f)
        self.option_add("*TCombobox*Listbox.font", self.f)
        self.option_add("*TCombobox*Listbox.selectBackground", YELLOW_SEL)
        self.option_add("*TCombobox*Listbox.selectForeground", "#000000")
        s = ttk.Style(self)
        s.theme_use("default")
        s.configure(".", background=FACE, font=self.f)
        s.configure("Treeview", background=LIGHT, fieldbackground=LIGHT, foreground="#000000",
                    rowheight=self.f.metrics("linespace") + 6, borderwidth=0)
        s.configure("Treeview.Heading", background=FACE, relief="raised", borderwidth=2, font=self.f)
        s.map("Treeview", background=[("selected", YELLOW_SEL)], foreground=[("selected", "#000000")])
        s.map("Treeview.Heading", background=[("active", FACE)])
        s.configure("TCombobox", fieldbackground=LIGHT, background=FACE, arrowsize=14, padding=2)
        s.map("TCombobox", fieldbackground=[("readonly", LIGHT)], selectbackground=[("readonly", LIGHT)],
              selectforeground=[("readonly", "#000000")])
        s.configure("TScrollbar", background=FACE, troughcolor="#E8E6E0", arrowsize=15)

    def build_ui(self):
        self.header = tk.Canvas(self, height=self.f_bold.metrics("linespace") + 8, bg=TITLE_FROM, highlightthickness=0, bd=0)
        self.header.pack(fill="x")
        self.header.bind("<Configure>", lambda e: self.draw_header())
        self.build_menu()
        self.build_toolbar()
        tk.Frame(self, height=2, bd=1, relief="groove", bg=FACE).pack(fill="x")
        self.build_statusbar()
        self.winbar = tk.Frame(self, bg=FACE, bd=1, relief="groove")
        self.winbar.pack(side="bottom", fill="x")
        self.mdi = tk.Frame(self, bg=MDI_BG, bd=2, relief="sunken")
        self.mdi.pack(fill="both", expand=True)

    def draw_header(self):
        c = self.header
        c.delete("all")
        w, h = c.winfo_width(), int(c["height"])
        for x in range(0, w, 2):
            c.create_rectangle(x, 0, x + 2, h, fill=blend(TITLE_FROM, TITLE_TO, x / max(1, w)), width=0)
        draw_icon(c, "app", 4, (h - 16) // 2)
        c.create_text(26, h // 2, text="БухКальк — Бухгалтерский калькулятор", anchor="w", font=self.f_bold)

    def build_menu(self):
        style = dict(tearoff=0, bg=FACE, fg="#000000", activebackground=YELLOW_SEL, activeforeground="#000000",
                     relief="raised", bd=2, font=self.f)
        bar = tk.Frame(self, bg=FACE)
        bar.pack(fill="x")
        tk.Frame(self, height=2, bd=1, relief="groove", bg=FACE).pack(fill="x")

        def submenu(title, items):
            mb = tk.Menubutton(bar, text=title, underline=0, bg=FACE, activebackground=YELLOW_SEL,
                               activeforeground="#000000", relief="flat", bd=1, padx=7, pady=2)
            mb.pack(side="left")
            m = tk.Menu(mb, **style)
            mb["menu"] = m
            for item in items:
                if item is None:
                    m.add_separator()
                else:
                    label, cmd, accel = item
                    m.add_command(label=label, command=cmd, accelerator=accel)
            return m

        page = lambda k: (lambda: self.show_page(k))
        submenu("Файл", [("Выгрузка отчёта...", page("reports"), ""), None, ("Выход", self.destroy, "")])
        submenu("Операции", [("Обзор", page("overview"), ""), ("Журнал операций", page("entries"), "Ctrl+J"),
                             ("Новая операция...", page("new_entry"), "Ctrl+N")])
        submenu("Расчёты", [("НДС-калькулятор...", page("calculator"), "Ctrl+1"),
                            ("Расчёт налога УСН...", page("usn"), "Ctrl+2"), None,
                            ("Больничный лист...", page("sick"), "Ctrl+3"),
                            ("Декретные...", page("maternity"), "Ctrl+4"),
                            ("Отпускные...", page("vacation"), "Ctrl+5")])
        submenu("Отчёты", [("Выгрузка отчёта...", page("reports"), "")])
        self.windows_menu = submenu("Окна", [])
        self.windows_menu.configure(postcommand=self.fill_windows_menu)
        submenu("Помощь", [("О программе", self.about, "F1")])

    def fill_windows_menu(self):
        m = self.windows_menu
        m.delete(0, "end")
        m.add_command(label="Закрыть все окна", command=self.close_all)
        if self.windows:
            m.add_separator()
        for i, win in enumerate(self.windows.values(), start=1):
            m.add_command(label=f"{i} {win.title}", command=lambda w=win: self.activate(w))

    def tool_button(self, parent, icon, tip, command, idle="flat"):
        cv = tk.Canvas(parent, width=24, height=22, bg=FACE, highlightthickness=0, bd=1, relief=idle)
        cv.pressed = False
        draw_icon(cv, icon, 5, 4)
        cv.bind("<Enter>", lambda e: (self.status_var.set(tip), cv.config(relief="sunken" if cv.pressed else "raised")))
        cv.bind("<Leave>", lambda e: (self.status_var.set(self.default_status), cv.config(relief="sunken" if cv.pressed else idle)))
        cv.idle = idle
        cv.bind("<Button-1>", lambda e: command())
        return cv

    def build_toolbar(self):
        tb = tk.Frame(self, bg=FACE)
        tb.pack(fill="x", pady=2)
        for _ in range(2):
            tk.Frame(tb, width=2, height=22, bd=1, relief="groove", bg=FACE).pack(side="left", padx=(3, 0))
        self.tools = {}
        items = [("new", "Новая операция (Ctrl+N)", "new_entry"), ("save", "Выгрузка отчёта", "reports"), None,
                 ("overview", "Обзор", "overview"), ("journal", "Журнал операций (Ctrl+J)", "entries"), None,
                 ("calc", "НДС-калькулятор (Ctrl+1)", "calculator"), ("usn", "Расчёт налога УСН (Ctrl+2)", "usn"),
                 ("sick", "Больничный лист (Ctrl+3)", "sick"), ("maternity", "Декретные (Ctrl+4)", "maternity"),
                 ("vacation", "Отпускные (Ctrl+5)", "vacation"), None, ("help", "О программе (F1)", None)]
        for item in items:
            if item is None:
                tk.Frame(tb, width=2, height=20, bd=1, relief="groove", bg=FACE).pack(side="left", padx=4)
                continue
            icon, tip, key = item
            cmd = (lambda k=key: self.show_page(k)) if key else self.about
            cv = self.tool_button(tb, icon, tip, cmd)
            cv.pack(side="left", padx=1)
            if key:
                self.tools[key] = cv

    def build_statusbar(self):
        sb = tk.Frame(self, bg=FACE)
        sb.pack(side="bottom", fill="x", pady=(1, 1))
        self.default_status = "Для получения подсказки нажмите F1"
        self.status_var = tk.StringVar(value=self.default_status)
        cell = dict(bg=FACE, bd=1, relief="sunken", anchor="w", padx=4)
        tk.Label(sb, textvariable=self.status_var, **cell).pack(side="left", fill="x", expand=True, padx=(2, 1))
        tk.Label(sb, text="NUM", **cell).pack(side="left", padx=1)
        tk.Label(sb, text="БД: ~/.buh_calculator", **cell).pack(side="left", padx=1)
        tk.Label(sb, text=date.today().strftime("%d.%m.%Y"), width=11, **{**cell, "anchor": "center"}).pack(side="left", padx=(1, 2))

    def refresh_winbar(self):
        for w in self.winbar.winfo_children():
            w.destroy()
        for win in self.windows.values():
            act = win is self.active_win
            text = win.title if len(win.title) < 30 else win.title[:29] + "…"
            tk.Button(self.winbar, text=text, relief="sunken" if act else "raised", bd=2,
                      bg=YELLOW_SOFT if act else FACE, activebackground=YELLOW_SEL, font=self.f_bold if act else self.f, padx=6, pady=0,
                      command=lambda w=win: self.activate(w)).pack(side="left", padx=2, pady=2)
        if not self.windows:
            tk.Label(self.winbar, text=" ", bg=FACE).pack(side="left", pady=3)

    def show_page(self, key):
        if key in self.windows:
            self.activate(self.windows[key])
            return
        title, icon, builder, resizable = self.pages[key]
        win = ChildWindow(self, key, title, icon, builder, resizable)
        n = len(self.windows)
        if key == "new_entry":
            win.move_to(180, 70)
        else:
            win.move_to(12 + (n % 8) * 26, 10 + (n % 8) * 24)
        self.windows[key] = win
        self.activate(win)

    def activate(self, win):
        if self.active_win is win and win.active:
            return
        for w in self.windows.values():
            if w.active:
                w.active = False
                w.draw_bar()
        win.active = True
        win.lift()
        win.draw_bar()
        self.active_win = win
        for key, cv in self.tools.items():
            self.set_pressed(cv, key == win.key)
        self.refresh_winbar()

    def set_pressed(self, cv, on):
        cv.pressed = on
        cv.config(relief="sunken" if on else cv.idle, bg=YELLOW_SOFT if on else FACE)

    def close_window(self, win):
        self.windows.pop(win.key, None)
        win.destroy()
        if self.active_win is win:
            self.active_win = None
            for cv in self.tools.values():
                self.set_pressed(cv, False)
            if self.windows:
                self.activate(list(self.windows.values())[-1])
        self.refresh_winbar()

    def close_all(self):
        for win in list(self.windows.values()):
            self.close_window(win)

    def on_global_click(self, e):
        w = e.widget
        while w is not None and not isinstance(w, ChildWindow):
            w = getattr(w, "master", None)
        if w is not None and self.windows.get(w.key) is w and w is not self.active_win:
            self.activate(w)

    def refresh_data(self):
        for key in ("overview", "entries"):
            if key in self.windows:
                self.windows[key].rebuild()

    def about(self):
        messagebox.showinfo("О программе", f"БухКальк — бухгалтерский калькулятор\n\nБаза данных:\n{DB_PATH}")

    def lbl(self, parent, text="", **kw):
        kw.setdefault("bg", FACE)
        return tk.Label(parent, text=text, **kw)

    def entry(self, parent, var, width=20, justify="left"):
        return tk.Entry(parent, textvariable=var, width=width, justify=justify, relief="sunken", bd=2,
                        bg=LIGHT, highlightthickness=0, insertwidth=1, selectbackground=YELLOW_SEL, selectforeground="#000000")

    def ro_entry(self, parent, var, width=18, color="#000000", font=None, bg=LIGHT):
        return tk.Entry(parent, textvariable=var, width=width, justify="right", relief="sunken", bd=2,
                        state="readonly", readonlybackground=bg, fg=color, highlightthickness=0,
                        font=font or self.f)

    def combo(self, parent, var, values, width=18):
        return ttk.Combobox(parent, textvariable=var, values=values, state="readonly", width=width, font=self.f)

    def radio(self, parent, text, var, value):
        return tk.Radiobutton(parent, text=text, variable=var, value=value, bg=FACE, activebackground=FACE,
                              selectcolor=LIGHT, highlightthickness=0)

    def group(self, parent, title):
        return tk.LabelFrame(parent, text=title, bg=FACE, relief="groove", bd=2, padx=10, pady=6)

    def button(self, parent, text, command, default=False, width=12):
        holder = tk.Frame(parent, bg="#000000" if default else FACE, padx=1, pady=1)
        tk.Button(holder, text=text, command=command, width=width, relief="raised", bd=2, bg=FACE,
                  activebackground=YELLOW_SOFT, highlightthickness=0, pady=1).pack()
        return holder

    def date_field(self, parent, var):
        f = tk.Frame(parent, bg=FACE)
        self.entry(f, var, width=12).pack(side="left")
        tk.Button(f, text="...", bd=2, relief="raised", bg=FACE, activebackground=FACE, padx=3, pady=0,
                  font=self.f_bold, command=lambda: var.set(date.today().isoformat())).pack(side="left", padx=(2, 0), fill="y")
        return f

    def button_row(self, body, buttons, note=None):
        tk.Frame(body, height=2, bd=1, relief="groove", bg=FACE).pack(fill="x", pady=(10, 8))
        row = tk.Frame(body, bg=FACE)
        row.pack(fill="x")
        if note:
            self.lbl(row, note, fg=GRAY).pack(side="left")
        for text, cmd, default in reversed(buttons):
            self.button(row, text, cmd, default).pack(side="right", padx=(6, 0))

    def note(self, body, text):
        box = tk.Frame(body, bg=INFO, bd=1, relief="solid")
        box.pack(fill="x", pady=(10, 0))
        cv = tk.Canvas(box, width=18, height=18, bg=INFO, highlightthickness=0)
        draw_icon(cv, "help", 1, 1)
        cv.pack(side="left", padx=(6, 4), pady=6)
        self.lbl(box, text, bg=INFO, wraplength=440, justify="left").pack(side="left", pady=6, padx=(0, 6))

    def table(self, parent, rows, height):
        frame = tk.Frame(parent, bd=2, relief="sunken", bg=LIGHT)
        cols = [("Дата", 92, "w"), ("Тип", 82, "w"), ("Категория", 100, "w"), ("Описание", 270, "w"),
                ("Сумма", 110, "e"), ("НДС", 56, "e")]
        tree = ttk.Treeview(frame, columns=[c[0] for c in cols], show="headings", height=height, selectmode="browse")
        for name, width, anchor in cols:
            tree.heading(name, text=name, anchor=anchor)
            tree.column(name, width=width, anchor=anchor, stretch=name == "Описание")
        sb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=sb.set)
        tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        for r in rows:
            kind = ("+ " if r[2] == "Доход" else "− ") + r[2]
            tree.insert("", "end", iid=str(r[0]), values=(r[1], kind, r[3], r[4] or "—", money(r[5]), f"{r[6]:g}%"))
        if not rows:
            tree.insert("", "end", values=("", "", "", "Пока нет операций. Операции → Новая операция", "", ""))
        else:
            tree.selection_set(str(rows[0][0]))
        return frame, tree

    def summary(self, rows):
        income = sum(r[5] for r in rows if r[2] == "Доход")
        expense = sum(r[5] for r in rows if r[2] == "Расход")
        vat = sum(r[5] * r[6] / 100 for r in rows)
        return income, expense, income-expense, vat

    def page_overview(self, body, win):
        rows = self.db.all()
        income, expense, profit, vat = self.summary(rows)
        g = self.group(body, "Итоги по всем операциям")
        g.pack(fill="x")
        values = [("Доходы:", income, GREEN), ("Расходы:", expense, RED),
                  ("Результат:", profit, NAVY), ("НДС по ставкам:", vat, "#806000")]
        for i, (title, val, col) in enumerate(values):
            pad = (0 if i == 0 else 14, 0)
            self.lbl(g, title).grid(row=0, column=i, sticky="w", padx=pad)
            self.ro_entry(g, tk.StringVar(value=money(val)), 16, col, self.f_total).grid(row=1, column=i, sticky="w", padx=pad, pady=(2, 0))
            self.lbl(g, "руб.", fg=GRAY).grid(row=2, column=i, sticky="w", padx=pad)
        self.lbl(body, "Последние операции:", font=self.f_bold).pack(anchor="w", pady=(10, 4))
        frame, _ = self.table(body, rows[:8], 8)
        frame.pack(fill="both", expand=True)
        self.button_row(body, [("Обновить", win.rebuild, False), ("Журнал...", lambda: self.show_page("entries"), False),
                               ("Закрыть", win.close, False)], note=f"Всего операций в базе: {len(rows)}")

    def page_entries(self, body, win):
        rows = self.db.all()
        tb = tk.Frame(body, bg=FACE)
        tb.pack(fill="x", pady=(0, 4))
        for icon, tip, cmd in [("new", "Новая операция", lambda: self.show_page("new_entry")),
                               ("delete", "Удалить выделенную операцию (Del)", lambda: self.delete_selected(win)),
                               ("refresh", "Обновить", win.rebuild)]:
            self.tool_button(tb, icon, tip, cmd, idle="raised").pack(side="left", padx=1)
        self.lbl(tb, "Интервал: все операции").pack(side="left", padx=10)
        frame, win.tree = self.table(body, rows, 16)
        frame.pack(fill="both", expand=True)
        win.tree.bind("<Delete>", lambda e: self.delete_selected(win))
        income, expense, _, _ = self.summary(rows)
        self.lbl(body, f"Операций: {len(rows)}    Доходы: {money(income)}    Расходы: {money(expense)}",
                 fg=GRAY).pack(anchor="w", pady=(6, 0))

    def delete_selected(self, win):
        sel = [i for i in win.tree.selection() if i.isdigit()]
        if not sel:
            messagebox.showinfo("Удаление", "Выделите операцию в журнале.")
            return
        values = win.tree.item(sel[0], "values")
        if messagebox.askyesno("Удаление", f"Удалить операцию от {values[0]} на сумму {values[4]}?"):
            self.db.delete(int(sel[0]))
            self.refresh_data()

    def page_new_entry(self, body, win):
        v = win.vars = {
            "date": tk.StringVar(value=date.today().isoformat()), "kind": tk.StringVar(value="Доход"),
            "category": tk.StringVar(value="Услуги"), "desc": tk.StringVar(), "amount": tk.StringVar(),
            "vat": tk.StringVar(value="0%"), "vat_hint": tk.StringVar()}
        form = tk.Frame(body, bg=FACE)
        form.pack(fill="x")

        def row(r, text):
            self.lbl(form, text).grid(row=r, column=0, sticky="w", pady=4, padx=(0, 20))

        row(0, "Дата:")
        df = tk.Frame(form, bg=FACE)
        self.date_field(df, v["date"]).pack(side="left")
        self.lbl(df, "(ГГГГ-ММ-ДД)", fg=GRAY).pack(side="left", padx=8)
        df.grid(row=0, column=1, sticky="w")
        row(1, "Тип операции:")
        kf = tk.Frame(form, bg=FACE)
        self.radio(kf, "Доход", v["kind"], "Доход").pack(side="left")
        self.radio(kf, "Расход", v["kind"], "Расход").pack(side="left", padx=16)
        kf.grid(row=1, column=1, sticky="w")
        row(2, "Категория:")
        self.combo(form, v["category"], CATEGORIES, 22).grid(row=2, column=1, sticky="w")
        row(3, "Описание:")
        self.entry(form, v["desc"], 46).grid(row=3, column=1, sticky="w")
        row(4, "Сумма:")
        af = tk.Frame(form, bg=FACE)
        self.entry(af, v["amount"], 16, "right").pack(side="left")
        self.lbl(af, "руб.").pack(side="left", padx=6)
        af.grid(row=4, column=1, sticky="w")
        row(5, "Ставка НДС:")
        vf = tk.Frame(form, bg=FACE)
        self.combo(vf, v["vat"], VAT_RATES, 6).pack(side="left")
        self.lbl(vf, textvariable=v["vat_hint"], fg=GRAY).pack(side="left", padx=10)
        vf.grid(row=5, column=1, sticky="w")

        def update_hint(*_):
            try:
                amount = float(v["amount"].get().replace(" ", "").replace(",", "."))
                v["vat_hint"].set(f"НДС по ставке: {money(amount * float(v['vat'].get().rstrip('%')) / 100)}")
            except ValueError:
                v["vat_hint"].set("")
        v["amount"].trace_add("write", update_hint)
        v["vat"].trace_add("write", update_hint)
        win.calc = lambda: self.add_entry(win, close=True)
        self.button_row(body, [("OK", lambda: self.add_entry(win, close=True), True),
                               ("Записать", lambda: self.add_entry(win, close=False), False),
                               ("Закрыть", win.close, False)])

    def add_entry(self, win, close):
        v = win.vars
        try:
            datetime.strptime(v["date"].get().strip(), "%Y-%m-%d")
            amount = float(v["amount"].get().replace(" ", "").replace(",", "."))
            if amount <= 0: raise ValueError
            vat = float(v["vat"].get().replace("%", ""))
            self.db.add(v["date"].get().strip(), v["kind"].get(), v["category"].get(),
                        v["desc"].get().strip(), amount, vat)
        except ValueError:
            messagebox.showerror("Проверьте данные", "Введите корректную дату в формате ГГГГ-ММ-ДД и сумму больше нуля.")
            return
        self.refresh_data()
        self.status_var.set("Операция записана")
        if close:
            win.close()
        else:
            v["amount"].set(""); v["desc"].set("")

    def calc_window(self, body, win, fields, results, callback, note):
        win.vars = {}
        g = self.group(body, "Исходные данные")
        g.pack(fill="x")
        g.grid_columnconfigure(0, minsize=LABEL_COL)
        for r, (key, label, default, kind, options) in enumerate(fields):
            var = win.vars[key] = tk.StringVar(value=default)
            self.lbl(g, label).grid(row=r, column=0, sticky="w", pady=3)
            if kind == "choice":
                w = self.combo(g, var, options, 22)
            elif kind == "radio":
                w = tk.Frame(g, bg=FACE)
                for opt in options:
                    self.radio(w, opt, var, opt).pack(side="left", padx=(0, 12))
            elif kind == "date":
                w = self.date_field(g, var)
            else:
                w = self.entry(g, var, 18, "right")
            w.grid(row=r, column=1, sticky="w", pady=3)
        rg = self.group(body, "Результат расчёта")
        rg.pack(fill="x", pady=(10, 0))
        rg.grid_columnconfigure(0, minsize=LABEL_COL)
        win.results = []
        for r, (label, unit) in enumerate(results):
            last = r == len(results) - 1
            lab = self.lbl(rg, label, font=self.f_bold if last else self.f)
            lab.grid(row=r, column=0, sticky="w", pady=3)
            var = tk.StringVar()
            self.ro_entry(rg, var, 20, NAVY if last else "#000000", self.f_bold if last else self.f,
                          YELLOW_SOFT if last else LIGHT).grid(row=r, column=1, sticky="w", pady=3)
            unit_lbl = self.lbl(rg, unit)
            unit_lbl.grid(row=r, column=2, sticky="w", padx=6)
            win.results.append((lab, var, unit_lbl))

        def run():
            try:
                out = callback(win)
            except (ValueError, KeyError) as e:
                messagebox.showerror("Ошибка", str(e) or "Проверьте введённые значения.")
                return
            for (lab, var, unit_lbl), (text, value, unit) in zip(win.results, out):
                lab.config(text=text); var.set(value); unit_lbl.config(text=unit)
            self.status_var.set("Расчёт выполнен")
        win.calc = run
        self.note(body, note)
        self.button_row(body, [("Рассчитать", run, True), ("Закрыть", win.close, False)])

    def cv(self, win, key):
        return win.vars[key].get().strip()

    def num(self, win, key):
        return float(self.cv(win, key).replace(" ", "").replace(",", "."))

    def page_calculator(self, body, win):
        self.calc_window(body, win,
            [("amount", "Исходная сумма, руб.:", "10000", "entry", None),
             ("rate", "Ставка НДС:", "20%", "choice", VAT_RATES),
             ("mode", "Режим расчёта:", "Начислить НДС", "radio", ["Начислить НДС", "Выделить НДС из суммы"])],
            [("НДС:", "руб."), ("Итог с НДС:", "руб.")],
            self.calculate, "Расчёт вспомогательный; налоговые правила зависят от вашей юрисдикции.")

    def calculate(self, win):
        try:
            amount = self.num(win, "amount")
            rate = float(self.cv(win, "rate").replace("%", "")) / 100
            if amount < 0: raise ValueError
        except ValueError:
            raise ValueError("Введите корректную числовую сумму.")
        if self.cv(win, "mode") == "Начислить НДС":
            vat = amount * rate
            return [("НДС:", money(vat), "руб."), ("Итог с НДС:", money(amount + vat), "руб.")]
        vat = amount - amount / (1 + rate) if rate else 0
        return [("НДС внутри суммы:", money(vat), "руб."), ("Сумма без НДС:", money(amount - vat), "руб.")]

    def page_usn(self, body, win):
        self.calc_window(body, win,
            [("type", "Тип налогоплательщика:", "ИП с работниками", "choice", ["ИП с работниками", "ИП без работников", "ООО"]),
             ("object", "Объект налогообложения:", "Доходы", "radio", ["Доходы", "Доходы минус расходы"]),
             ("period", "Расчётный период:", "Год", "choice", ["Первый квартал", "Полугодие", "9 месяцев", "Год"]),
             ("income", "Доходы за период, руб.:", "1000000", "entry", None),
             ("expenses", "Расходы за период, руб.:", "300000", "entry", None),
             ("rate", "Ставка налога, %:", "6", "entry", None),
             ("contrib", "Страховые взносы, руб.:", "0", "entry", None)],
            [("Налоговая база:", "руб."), ("Налог до уменьшения:", "руб."), ("К уплате после взносов:", "руб.")],
            self.calculate_usn, "Уменьшение на взносы: до 50% для работодателей, до 100% для ИП без работников.")

    def calculate_usn(self, win):
        error = "Проверьте числовые значения и ставку от 0 до 100."
        try:
            income, expenses, rate, contrib = (self.num(win, k) for k in ("income", "expenses", "rate", "contrib"))
        except ValueError:
            raise ValueError(error)
        if min(income, expenses, rate, contrib) < 0 or rate > 100: raise ValueError(error)
        taxable = income if self.cv(win, "object") == "Доходы" else max(0, income-expenses)
        tax = taxable*rate/100
        # Typical simplified reduction cap: 50% for employers/companies; potentially 100% for sole proprietors without staff.
        cap = 1.0 if self.cv(win, "type") == "ИП без работников" else 0.5
        due = max(0, tax-min(contrib, tax*cap)) if self.cv(win, "object") == "Доходы" else tax
        return [("Налоговая база:", money(taxable), "руб."), ("Налог до уменьшения:", money(tax), "руб."),
                ("К уплате после взносов:", money(due), "руб.")]

    def page_sick(self, body, win):
        self.calc_window(body, win,
            [("start", "Дата начала:", date.today().isoformat(), "date", None),
             ("end", "Дата окончания:", date.today().isoformat(), "date", None),
             ("income", "Заработок за расчётные 2 года, руб.:", "1200000", "entry", None),
             ("experience", "Страховой стаж:", "5–8 лет", "choice", ["Менее 5 лет", "5–8 лет", "8 лет и более"])],
            [("Дней нетрудоспособности:", ""), ("Средний дневной заработок:", "руб."), ("Ориентировочное пособие:", "руб.")],
            self.calculate_sick, "Расчёт ориентировочный. Перед применением проверьте действующие правила для вашего региона.")

    def calculate_sick(self, win):
        error = "Проверьте даты в формате ГГГГ-ММ-ДД и заработок."
        try:
            start = datetime.strptime(self.cv(win, "start"), "%Y-%m-%d").date()
            end = datetime.strptime(self.cv(win, "end"), "%Y-%m-%d").date()
            income = self.num(win, "income")
        except ValueError:
            raise ValueError(error)
        days = (end-start).days+1
        if days < 1 or income < 0: raise ValueError(error)
        factor = {"Менее 5 лет": 0.6, "5–8 лет": 0.8, "8 лет и более": 1.0}[self.cv(win, "experience")]
        avg = income/730
        return [("Дней нетрудоспособности:", str(days), ""), ("Средний дневной заработок:", money(avg), "руб."),
                ("Ориентировочное пособие:", money(avg*factor*days), "руб.")]

    def page_maternity(self, body, win):
        self.calc_window(body, win,
            [("income", "Заработок за расчётные 2 года, руб.:", "1200000", "entry", None),
             ("days", "Длительность отпуска, дней:", "140", "choice", ["140", "156", "194"]),
             ("excluded", "Исключаемые дни:", "0", "entry", None),
             ("min_day", "Ср. дневной заработок, если известен:", "0", "entry", None)],
            [("Средний дневной заработок:", "руб."), ("Период отпуска, дней:", ""), ("Ориентировочное пособие:", "руб.")],
            self.calculate_maternity, "Расчёт ориентировочный. Перед применением проверьте действующие правила для вашего региона.")

    def calculate_maternity(self, win):
        error = "Проверьте сумму заработка и количество исключаемых дней."
        try:
            income, excluded, known = self.num(win, "income"), self.num(win, "excluded"), self.num(win, "min_day")
            days = int(self.cv(win, "days"))
        except ValueError:
            raise ValueError(error)
        denominator = 730-excluded
        if income < 0 or excluded < 0 or denominator <= 0 or known < 0: raise ValueError(error)
        avg = known if known > 0 else income/denominator
        return [("Средний дневной заработок:", money(avg), "руб."), ("Период отпуска, дней:", str(days), ""),
                ("Ориентировочное пособие:", money(avg*days), "руб.")]

    def page_vacation(self, body, win):
        self.calc_window(body, win,
            [("income", "Начисления за расчётный период, руб.:", "1200000", "entry", None),
             ("months", "Полных месяцев в периоде:", "12", "entry", None),
             ("days", "Дней отпуска:", "14", "entry", None),
             ("excluded", "Неучитываемые дни (оценка):", "0", "entry", None)],
            [("Средний дневной заработок:", "руб."), ("Дней отпуска:", ""), ("Ориентировочные отпускные:", "руб.")],
            self.calculate_vacation, "Средний дневной заработок = начисления / (месяцы × 29,3 − неучитываемые дни).")

    def calculate_vacation(self, win):
        error = "Проверьте начисления, месяцы, дни отпуска и исключаемые дни."
        try:
            income, months, days, excluded = (self.num(win, k) for k in ("income", "months", "days", "excluded"))
        except ValueError:
            raise ValueError(error)
        if min(income, months, days, excluded) < 0 or months <= 0: raise ValueError(error)
        calendar_days = max(1, months*29.3-excluded)
        avg = income/calendar_days
        return [("Средний дневной заработок:", money(avg), "руб."), ("Дней отпуска:", f"{days:g}", ""),
                ("Ориентировочные отпускные:", money(avg*days), "руб.")]

    def page_reports(self, body, win):
        win.vars = {"start": tk.StringVar(), "end": tk.StringVar(), "fmt": tk.StringVar(value="xlsx")}
        g = self.group(body, "Период отчёта")
        g.pack(fill="x")
        self.lbl(g, "с:").grid(row=0, column=0, sticky="w")
        self.date_field(g, win.vars["start"]).grid(row=0, column=1, sticky="w", padx=(6, 20))
        self.lbl(g, "по:").grid(row=0, column=2, sticky="w")
        self.date_field(g, win.vars["end"]).grid(row=0, column=3, sticky="w", padx=6)
        self.lbl(g, "Оставьте обе даты пустыми, чтобы выгрузить все операции.", fg=GRAY).grid(
            row=1, column=0, columnspan=4, sticky="w", pady=(6, 0))
        fg = self.group(body, "Формат выгрузки")
        fg.pack(fill="x", pady=(10, 0))
        self.radio(fg, "CSV (разделитель «;», UTF-8)", win.vars["fmt"], "csv").pack(anchor="w")
        self.radio(fg, "Excel (.xlsx) с листом «Сводка»", win.vars["fmt"], "xlsx").pack(anchor="w")
        cg = self.group(body, "Колонки отчёта")
        cg.pack(fill="x", pady=(10, 0))
        self.lbl(cg, "ID, дата, тип, категория, описание, сумма, ставка НДС, сумма НДС.").pack(anchor="w")
        win.calc = lambda: self.export(win)
        self.button_row(body, [("Выгрузить", lambda: self.export(win), True), ("Закрыть", win.close, False)])

    def export(self, win):
        fmt = win.vars["fmt"].get()
        start, end = win.vars["start"].get().strip(), win.vars["end"].get().strip()
        try:
            if bool(start) != bool(end): raise ValueError
            if start:
                datetime.strptime(start, "%Y-%m-%d"); datetime.strptime(end, "%Y-%m-%d")
                if start > end: raise ValueError
            rows = self.db.all(start or None, end or None)
        except ValueError:
            messagebox.showerror("Период не задан корректно", "Укажите обе даты в формате ГГГГ-ММ-ДД или оставьте обе пустыми.")
            return
        if not rows:
            messagebox.showinfo("Нет данных", "За выбранный период операций нет.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".csv" if fmt == "csv" else ".xlsx",
                  filetypes=[("CSV", "*.csv")] if fmt == "csv" else [("Excel", "*.xlsx")],
                  initialfile=f"buh_report_{date.today().isoformat()}")
        if not path: return
        headers = ["ID", "Дата", "Тип", "Категория", "Описание", "Сумма", "Ставка НДС, %", "Сумма НДС"]
        data = [[r[0], r[1], r[2], r[3], r[4] or "", r[5], r[6], round(r[5]*r[6]/100, 2)] for r in rows]
        try:
            if fmt == "csv":
                with open(path, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f, delimiter=";"); writer.writerow(headers); writer.writerows(data)
            else:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill
                wb = Workbook(); ws = wb.active; ws.title = "Операции"
                ws.append(headers)
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="000000")
                    cell.fill = PatternFill("solid", fgColor=TITLE_FROM.replace("#", ""))
                for row in data: ws.append(row)
                ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
                for col, width in zip("ABCDEFGH", [8, 14, 14, 18, 32, 16, 16, 16]): ws.column_dimensions[col].width = width
                ws2 = wb.create_sheet("Сводка")
                inc, exp, prof, vat = self.summary(rows)
                ws2.append(["Показатель", "Значение"])
                for item in [("Доходы", inc), ("Расходы", exp), ("Результат", prof), ("НДС по операциям", vat)]:
                    ws2.append(list(item))
                ws2.column_dimensions["A"].width = 28; ws2.column_dimensions["B"].width = 20
                wb.save(path)
            messagebox.showinfo("Готово", f"Отчёт сохранён:\n{path}")
            self.status_var.set(f"Отчёт сохранён: {path}")
        except ImportError:
            messagebox.showerror("Не хватает библиотеки", "Для Excel установите openpyxl: pip install openpyxl")
        except Exception as e:
            messagebox.showerror("Ошибка экспорта", str(e))

if __name__ == "__main__":
    app = App()
    app.mainloop()
